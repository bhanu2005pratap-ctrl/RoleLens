"""
Turn an uploaded resume file (pdf / docx / txt) into plain text.
Kept dependency-light on purpose: pypdf + python-docx only, no OCR, no
local ML models -- this runs inside a Vercel serverless function.
"""
from __future__ import annotations
import io
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
import docx


async def extract_text(file: UploadFile) -> str:
    raw = await file.read()
    name = (file.filename or "").lower()

    # PDF signature
    if raw.startswith(b"%PDF-"):
        return _extract_pdf(raw)

    # DOCX is a ZIP container
    if raw.startswith(b"PK\x03\x04") and name.endswith(".docx"):
        return _extract_docx(raw)

    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="ignore")

    raise HTTPException(
        status_code=400,
        detail="Unsupported or invalid resume file. Please upload a valid .pdf, .docx, or .txt file.",
    )


def _extract_pdf(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="Couldn't extract text from this PDF -- it may be a scanned image. "
                   "Try a text-based PDF or a .docx export instead.",
        )
    return text


def _extract_docx(raw: bytes) -> str:
    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(p for p in parts if p.strip())
    if not text:
        raise HTTPException(status_code=422, detail="This .docx file appears to have no extractable text.")
    return text


def clip(text: str, max_chars: int = 12000) -> str:
    """Guard against pathologically long documents blowing up token/latency budgets."""
    return text if len(text) <= max_chars else text[:max_chars]
